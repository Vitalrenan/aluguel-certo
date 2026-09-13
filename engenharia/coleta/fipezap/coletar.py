"""
FipeZAP ingestion -- the published series into the raw layer.

    python -m coleta.fipezap.coletar                    # baixa, valida, relata
    python -m coleta.fipezap.coletar --write
    python -m coleta.fipezap.coletar --arquivo x.xlsx --write

WHAT THIS IS, AND WHAT IT IS NOT. Not scraping. FipeZAP publishes a historical
series as an `.xlsx` at a stable address; this fetches that file once per run
and tabulates it. No pagination, no search, no listing page, and no personal
data anywhere -- it is an aggregate market index.

WHAT CHANGED FROM THE OLD COLLECTOR. It wrote straight into the processed
layer, and on the way it kept the LAST month of three fields per city and threw
the rest away -- before anything reached disk. That is a modelling decision
taken at ingestion, which is where it must not live: the day the calculator
needs the price series rather than one month of it, the whole download happens
again.

Here the raw layer keeps the WHOLE series: every city, every month, every
tracked field, in long form. The cut lives in `tratamento.fipezap.derivar`.

THE LAYOUT IS CHECKED AGAINST LABELS, NOT TRUSTED BY POSITION. The spreadsheet
belongs to someone else and can change shape without telling anyone. A moved
column that goes unchecked writes the sale price into the yield field, and every
number downstream stays plausible.
"""
from __future__ import annotations

import argparse
import sys
import tempfile
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

URL = "https://downloads.fipe.org.br/indices/fipezap/fipezap-serieshistoricas.xlsx"
RAW_PADRAO = "dados/01_raw/fipezap"

# Identifica quem somos e como falar conosco. NAO e UA de navegador: o CLAUDE.md
# proibe UA falso, que e evasao. Isto e o contrario -- e assinatura.
AGENTE = ("AluguelCerto/1.0 (+https://github.com/aluguelcerto; "
          "ingestao de indice publico)")

COL_DATA = 1
LINHA_CABECALHO = (2, 3, 4)          # 1-based: bloco, sub-bloco, coluna

# Posicoes medidas em 2026-09-09, conferidas contra os rotulos antes do uso.
BLOCOS = {
    "venda_m2": (17, "Venda", "Preço médio"),
    "venda_var12": (12, "Venda", "Var. em 12 meses"),
    "locacao_m2": (37, "Locação", "Preço médio"),
    "locacao_var12": (32, "Locação", "Var. em 12 meses"),
    "yield_mensal": (42, "Rentabilidade", "mensalizada"),
}


class ColetaInvalida(Exception):
    """The spreadsheet is not what we expect. Raises; never a crooked number."""


def baixa(url: str = URL, destino: Path | None = None) -> Path:
    """One request, a published file. Network failure raises."""
    destino = destino or Path(tempfile.mkdtemp(prefix="fipezap_")) / "serie.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers={"User-Agent": AGENTE})
    with urllib.request.urlopen(req, timeout=180) as r:  # noqa: S310
        if r.status != 200:
            raise ColetaInvalida(f"{url} respondeu {r.status}")
        tipo = r.headers.get("Content-Type", "")
        if "spreadsheet" not in tipo and "excel" not in tipo:
            # HTTP 200 com HTML e a falha silenciosa classica: pagina de erro
            # ou desafio servida com status de sucesso.
            raise ColetaInvalida(
                f"{url} devolveu 200 mas Content-Type {tipo!r} -- nao e planilha")
        destino.write_bytes(r.read())
    if destino.stat().st_size < 100_000:
        raise ColetaInvalida(
            f"planilha com {destino.stat().st_size} bytes -- pequena demais")
    return destino


def _ffill(linha) -> list[str]:
    """Merged cells carry a value only in the first column; propagate right."""
    saida, ultimo = [], ""
    for c in linha:
        if c not in (None, ""):
            ultimo = str(c)
        saida.append(ultimo)
    return saida


def confere_layout(ws, cidade: str) -> None:
    """
    Column positions are checked against the LABELS before being used.

    The difference between a run that stops and one that writes the sale price
    into the yield field.
    """
    linhas = [_ffill(r) for r in ws.iter_rows(
        min_row=min(LINHA_CABECALHO), max_row=max(LINHA_CABECALHO),
        values_only=True)]
    for campo, (col, bloco, rotulo) in BLOCOS.items():
        if col >= len(linhas[0]):
            raise ColetaInvalida(f"{cidade}: coluna {col} nao existe")
        texto = " ".join(l[col] for l in linhas).lower()
        for esperado in (bloco.lower(), rotulo.lower()):
            if esperado not in texto:
                raise ColetaInvalida(
                    f"{cidade}: coluna {col} devia ser {campo!r} "
                    f"({bloco}/{rotulo}) e o cabecalho diz {texto!r}. "
                    f"O layout mudou -- reveja BLOCOS antes de rodar.")


def serie_da_cidade(ws, cidade: str) -> list[dict]:
    """
    EVERY month of every tracked field, in long form.

    Long rather than wide on purpose: the fields do not end in the same month
    -- price runs to 2026-08 while yield stops at 2026-07 -- and a wide table
    forces a decision about what to do with the ragged edge. Long form has no
    edge: a month simply has the rows it has.
    """
    confere_layout(ws, cidade)
    linhas = []
    for r in ws.iter_rows(min_row=max(LINHA_CABECALHO) + 1, values_only=True):
        data = r[COL_DATA]
        if not isinstance(data, datetime):
            continue
        mes = data.strftime("%Y-%m")
        for campo, (i, _, _) in BLOCOS.items():
            if i < len(r) and isinstance(r[i], (int, float)):
                linhas.append({"cidade": cidade, "mes_referencia": mes,
                               "campo": campo, "valor": float(r[i])})
    return linhas


def tabula(caminho: Path) -> pd.DataFrame:
    """The whole workbook: one row per city, month and field."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    linhas, recusadas = [], {}
    for nome in wb.sheetnames:
        try:
            linhas.extend(serie_da_cidade(wb[nome], nome))
        except ColetaInvalida as exc:
            # Aba que nao e de cidade (indice, metodologia) cai aqui. Contada e
            # relatada, nunca silenciada: se TODAS cairem, a corrida falha.
            recusadas[nome] = str(exc)[:80]
    wb.close()

    if not linhas:
        raise ColetaInvalida(
            f"nenhuma aba tabulavel em {caminho.name}. "
            f"{len(recusadas)} recusada(s): {list(recusadas)[:5]}")
    d = pd.DataFrame(linhas)
    d.attrs["recusadas"] = recusadas
    return d


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--arquivo", default=None, help="usa um .xlsx local")
    ap.add_argument("--raw-dir", default=RAW_PADRAO)
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--guardar", action="store_true",
                    help="mantem o .xlsx baixado")
    args = ap.parse_args(argv)

    if args.arquivo:
        origem, temporario = Path(args.arquivo), False
    else:
        origem, temporario = baixa(), True

    try:
        d = tabula(origem)
    finally:
        # O .xlsx baixado e descartado. Nao por conter PII -- nao contem -- mas
        # porque a regra do projeto e que o arquivo de origem nao fica.
        if temporario and not args.guardar and origem.exists():
            origem.unlink()

    recusadas = d.attrs.get("recusadas", {})
    print("=" * 72)
    print("FipeZAP -- serie historica completa")
    print("=" * 72)
    print(f"  {len(d):,} linhas   {d.cidade.nunique()} cidades   "
          f"{d.campo.nunique()} campos")
    print(f"  meses de {d.mes_referencia.min()} a {d.mes_referencia.max()}")
    if recusadas:
        print(f"  {len(recusadas)} aba(s) nao tabulavel(is): "
              f"{', '.join(list(recusadas)[:6])}")

    print("\n  cobertura por campo:")
    for campo, sub in d.groupby("campo"):
        print(f"    {campo:<16}{len(sub):>8,} linhas   "
              f"{sub.cidade.nunique():>3} cidades   ate {sub.mes_referencia.max()}")

    alvo_cidades = ["Santos", "São Paulo", "Guarujá", "São Vicente", "Praia Grande"]
    presentes = [c for c in alvo_cidades if c in set(d.cidade)]
    print(f"\n  cidades da nossa base presentes: {len(presentes)}/{len(alvo_cidades)}"
          f"  {presentes}")
    faltando = [c for c in alvo_cidades if c not in presentes]
    if faltando:
        print(f"  AUSENTES: {faltando}")

    if not args.write:
        print("\nsem --write: nada foi gravado.")
        return 0

    agora = datetime.now(timezone.utc)
    pasta = Path(args.raw_dir) / agora.strftime("%Y") / agora.strftime("%m")
    pasta.mkdir(parents=True, exist_ok=True)
    alvo = pasta / "serie.parquet"
    try:
        d.to_parquet(alvo, index=False)
    except Exception as exc:
        raise ColetaInvalida(f"falha ao gravar {alvo}: {exc}") from exc
    print(f"\n  gravado: {alvo}  ({len(d):,} linhas)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
